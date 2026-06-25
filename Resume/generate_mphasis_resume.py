"""
Resume Generator - Navneet Vishwakarma
Tailored for: Mphasis Trainee Software Engineer - Systems (Job ID: 118377-1-10)
100% honest content. Aligned to Systems Engineer / Technical Support role.
"""

import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Navneet_Vishwakarma_Mphasis_Trainee_SE.docx")

DARK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x16, 0x53, 0x8D)
TEXT = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE_COLOR = "165390"
FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"


def remove_paragraph_spacing(paragraph):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(11.5)


def add_thin_rule(doc, color=RULE_COLOR, thickness="8000"):
    p = doc.add_paragraph()
    remove_paragraph_spacing(p)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{thickness}" w:space="1" w:color="{color}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def styled_run(paragraph, text, font_name=FONT_BODY, size=Pt(10), color=TEXT,
               bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    east_asia = rPr.find(qn('w:rFonts'))
    if east_asia is not None and east_asia.get(qn('w:eastAsiaTheme')) is not None:
        del east_asia.attrib[qn('w:eastAsiaTheme')]
    return run


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = Pt(13)
    styled_run(p, title.upper(), font_name=FONT_HEADING, size=Pt(10.5),
               color=ACCENT, bold=True)
    add_thin_rule(doc)


def add_bullet(doc, text, indent=Inches(0.25)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = Pt(11)
    p.paragraph_format.left_indent = indent
    p.paragraph_format.first_line_indent = -Inches(0.15)
    styled_run(p, "▪  ", size=Pt(7), color=ACCENT)
    styled_run(p, text, size=Pt(9))


def build_resume():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    style = doc.styles['Normal']
    style.font.name = FONT_BODY
    style.font.size = Pt(10)
    style.font.color.rgb = TEXT
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    # ═══════════════════════════════════════════════════════════
    # HEADER
    # ═══════════════════════════════════════════════════════════
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = Pt(22)
    styled_run(p_name, "NAVNEET VISHWAKARMA", font_name=FONT_HEADING,
               size=Pt(18), color=ACCENT, bold=True)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.line_spacing = Pt(12)
    styled_run(p_title,
               "Technical Support  |  System Operations  |  Troubleshooting  |  Documentation",
               size=Pt(8.5), color=MUTED)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(1)
    p_contact.paragraph_format.line_spacing = Pt(11)
    styled_run(p_contact, "+91 8435061006  |  Rajvl132011@gmail.com  |  Satna, Madhya Pradesh, India",
               size=Pt(8.5), color=TEXT)

    p_links = doc.add_paragraph()
    p_links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_links.paragraph_format.space_after = Pt(2)
    p_links.paragraph_format.line_spacing = Pt(11)
    styled_run(p_links, "LinkedIn: ", size=Pt(8.5), color=TEXT)
    styled_run(p_links, "linkedin.com/in/navneet1206", size=Pt(8.5), color=ACCENT)
    styled_run(p_links, "  |  GitHub: ", size=Pt(8.5), color=TEXT)
    styled_run(p_links, "github.com/Navneet1206", size=Pt(8.5), color=ACCENT)

    add_thin_rule(doc, thickness="12000")

    # ═══════════════════════════════════════════════════════════
    # PROFESSIONAL SUMMARY
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(2)
    p_sum.paragraph_format.space_after = Pt(2)
    p_sum.paragraph_format.line_spacing = Pt(12)
    styled_run(
        p_sum,
        "Computer Science Engineering student with experience in technical support, system "
        "operations, troubleshooting, documentation, and process management. Proficient in "
        "Linux environments, Microsoft Office, data analysis, and operational support "
        "activities. Strong problem-solving, communication, and collaboration skills with the "
        "ability to learn new technologies quickly and work effectively in team environments. "
        "Seeking opportunities as a Trainee Software Engineer, Systems Engineer, Technical "
        "Support Engineer, or Infrastructure Support Engineer.",
        size=Pt(9)
    )

    # ═══════════════════════════════════════════════════════════
    # TECHNICAL SKILLS
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Technical Skills")

    skills_data = [
        ("Operating Systems",
         "Linux, Windows"),
        ("Office & Productivity",
         "Microsoft Excel, Microsoft Word, Microsoft PowerPoint"),
        ("Technical Support",
         "Troubleshooting, Issue Resolution, Documentation, Process Monitoring"),
        ("Tools & Technologies",
         "Git, Basic Shell Scripting, Microsoft Office"),
        ("Soft Skills",
         "Communication, Team Collaboration, Problem Solving, SOP Compliance"),
    ]

    for label, value in skills_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = Pt(11.5)
        p.paragraph_format.left_indent = Inches(0.1)
        styled_run(p, f"{label}:  ", size=Pt(9), color=DARK, bold=True)
        styled_run(p, value, size=Pt(9))

    # ═══════════════════════════════════════════════════════════
    # PROFESSIONAL EXPERIENCE
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Professional Experience")

    p_job = doc.add_paragraph()
    p_job.paragraph_format.space_before = Pt(2)
    p_job.paragraph_format.space_after = Pt(0)
    p_job.paragraph_format.line_spacing = Pt(13)
    styled_run(p_job, "Operations Support Associate",
               size=Pt(10), color=DARK, bold=True)

    p_co = doc.add_paragraph()
    p_co.paragraph_format.space_before = Pt(0)
    p_co.paragraph_format.space_after = Pt(2)
    p_co.paragraph_format.line_spacing = Pt(11)
    styled_run(p_co, "Savayas Life and Balance Pvt. Ltd.",
               size=Pt(9), color=TEXT, italic=True)
    styled_run(p_co, "  |  March 2025 – Present  |  Satna, MP",
               size=Pt(8.5), color=MUTED)

    bullets_exp = [
        "Maintained operational records and reports using Microsoft Excel and Office tools.",
        "Assisted in daily process monitoring and administrative activities.",
        "Coordinated with internal teams to ensure smooth execution of operational tasks.",
        "Performed data verification, documentation, and record maintenance.",
        "Supported issue tracking and resolution through proper reporting procedures.",
        "Followed organizational SOPs and maintained process compliance.",
        "Prepared routine reports and updated operational data accurately.",
        "Assisted in handling customer and internal operational requests.",
    ]
    for b in bullets_exp:
        add_bullet(doc, b)

    # ═══════════════════════════════════════════════════════════
    # ACADEMIC PROJECTS
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Academic Projects")

    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_before = Pt(2)
    p_proj.paragraph_format.space_after = Pt(1)
    p_proj.paragraph_format.line_spacing = Pt(12)
    styled_run(p_proj, "Inventory Management System",
               size=Pt(10), color=DARK, bold=True)

    proj_bullets = [
        "Developed a basic inventory tracking solution for maintaining product records and "
        "stock information.",
        "Generated reports and organized data using spreadsheets and documentation tools.",
        "Improved accuracy of record keeping and data management processes.",
    ]
    for b in proj_bullets:
        add_bullet(doc, b)

    # ═══════════════════════════════════════════════════════════
    # CERTIFICATIONS
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Certifications")

    certs = [
        "Linux Administration Training (CX-501) – Codenixia",
        "Linux Server Hardening & Security Training (CX-701) – Codenixia",
        "SUSE Linux Enterprise Server (SLES) Administration Training",
    ]
    for c in certs:
        add_bullet(doc, c)

    # ═══════════════════════════════════════════════════════════
    # EDUCATION
    # ═══════════════════════════════════════════════════════════
    add_section_heading(doc, "Education")

    p_deg = doc.add_paragraph()
    p_deg.paragraph_format.space_before = Pt(2)
    p_deg.paragraph_format.space_after = Pt(0)
    p_deg.paragraph_format.line_spacing = Pt(13)
    styled_run(p_deg, "Bachelor of Technology (B.Tech) – Computer Science & Engineering",
               size=Pt(9.5), color=DARK, bold=True)

    p_uni = doc.add_paragraph()
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(1)
    p_uni.paragraph_format.line_spacing = Pt(11)
    styled_run(p_uni, "AKS University, Satna, Madhya Pradesh",
               size=Pt(9), color=TEXT, italic=True)
    styled_run(p_uni, "  |  2022 – 2026  |  CGPA: 8.29 / 10",
               size=Pt(8.5), color=MUTED)

    doc.save(OUTPUT_FILE)
    print(f"Resume saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_resume()
