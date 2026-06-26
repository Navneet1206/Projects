"""
Viatris Application - Resume (Single Page)
"""

import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Navneet_Vishwakarma_Resume.docx")

DARK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x16, 0x53, 0x8D)
TEXT = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x55, 0x55, 0x55)
HIGHLIGHT_BG = "E8F0FE"
RULE_COLOR = "165390"
FONT = "Calibri"


def remove_paragraph_spacing(p):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(10)


def add_thin_rule(doc, thickness="6000"):
    p = doc.add_paragraph()
    remove_paragraph_spacing(p)
    p.paragraph_format.space_after = Pt(1.5)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{thickness}" w:space="1" w:color="{RULE_COLOR}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def styled_run(paragraph, text, size=Pt(8.5), color=TEXT, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = size
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('w:rFonts'))
    if ea is not None and ea.get(qn('w:eastAsiaTheme')) is not None:
        del ea.attrib[qn('w:eastAsiaTheme')]
    return run


def section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = Pt(11)
    styled_run(p, title.upper(), size=Pt(9), color=ACCENT, bold=True)
    add_thin_rule(doc)


def bullet(doc, text, indent=Inches(0.2)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0.3)
    p.paragraph_format.space_after = Pt(0.3)
    p.paragraph_format.line_spacing = Pt(10)
    p.paragraph_format.left_indent = indent
    p.paragraph_format.first_line_indent = -Inches(0.13)
    styled_run(p, "▪ ", size=Pt(6), color=ACCENT)
    styled_run(p, text, size=Pt(8))


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Cm(0.7)
    section.bottom_margin = Cm(0.5)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(8.5)
    style.font.color.rgb = TEXT
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    # ═══════════════════════════════════════════════════════════
    # HEADER
    # ═══════════════════════════════════════════════════════════
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(1)
    p_name.paragraph_format.line_spacing = Pt(17)
    styled_run(p_name, "NAVNEET VISHWAKARMA", size=Pt(15), color=ACCENT, bold=True)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(1.5)
    p_title.paragraph_format.line_spacing = Pt(10)
    styled_run(p_title,
               "IT Support Engineer  |  Linux Administrator  |  Infrastructure Operations  |  Technical Support",
               size=Pt(8), color=MUTED)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(0.5)
    p_contact.paragraph_format.line_spacing = Pt(10)
    styled_run(p_contact,
               "+91-84350-61006  |  rajvl132011@gmail.com  |  Satna, MP, India (Willing to Relocate)",
               size=Pt(8))

    p_links = doc.add_paragraph()
    p_links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_links.paragraph_format.space_after = Pt(1)
    p_links.paragraph_format.line_spacing = Pt(10)
    styled_run(p_links, "LinkedIn: ", size=Pt(8))
    styled_run(p_links, "linkedin.com/in/navneet1206", size=Pt(8), color=ACCENT)
    styled_run(p_links, "  |  GitHub: ", size=Pt(8))
    styled_run(p_links, "github.com/navneetvishwakarma", size=Pt(8), color=ACCENT)

    # Referral highlight right below header
    p_ref_top = doc.add_paragraph()
    p_ref_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ref_top.paragraph_format.space_after = Pt(1.5)
    p_ref_top.paragraph_format.line_spacing = Pt(11)
    ref_run = p_ref_top.add_run("  Referred by: Mr. Chirag Sen  ")
    ref_run.font.name = FONT
    ref_run.font.size = Pt(8.5)
    ref_run.font.color.rgb = ACCENT
    ref_run.bold = True
    rPr_ref = ref_run._r.get_or_add_rPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{HIGHLIGHT_BG}"/>')
    rPr_ref.append(shading)

    add_thin_rule(doc, thickness="10000")

    # ═══════════════════════════════════════════════════════════
    # PROFESSIONAL SUMMARY
    # ═══════════════════════════════════════════════════════════
    section_heading(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(1)
    p_sum.paragraph_format.line_spacing = Pt(10)
    styled_run(
        p_sum,
        "IT Support & Linux System Administrator with 2+ years of experience in Linux Administration, "
        "Infrastructure Operations, and Technical Support. Skilled in managing SLES, Ubuntu, and "
        "Windows Server environments. Proficient in Bash Scripting, Ansible automation, Docker, "
        "system monitoring, log analysis, and Root Cause Analysis. Demonstrated ability to resolve "
        "P1/P2 incidents, implement CIS/NIST security baselines, and maintain 99.5% service availability.",
        size=Pt(8)
    )

    # ═══════════════════════════════════════════════════════════
    # TECHNICAL SKILLS
    # ═══════════════════════════════════════════════════════════
    section_heading(doc, "Technical Skills")

    skills = [
        ("Operating Systems", "SLES 12/15, Ubuntu Linux, openSUSE Leap, Windows Server 2016/2019"),
        ("Linux Admin", "User/Group Mgmt, File Systems, Package Mgmt (RPM, APT), systemd, Security Hardening, Performance Monitoring (top, htop, vmstat, iostat, sar)"),
        ("IT Support", "Incident Mgmt, RCA, Log Analysis, Ticketing Systems, HW/SW Troubleshooting, Remote & End-User Support"),
        ("Automation", "Bash Scripting, Ansible (Playbooks, Roles, Handlers, Inventory), Cron, Systemd Timers"),
        ("Virtualization & Containers", "VMware vSphere, KVM, Docker, Docker Compose"),
        ("Networking", "TCP/IP, DNS, DHCP, SSH, HTTP/HTTPS, Firewall (UFW, firewalld, iptables), VLANs, NAT"),
        ("Security", "CIS/NIST Hardening, Audit Logging, Lynis, PAM Config, SSH Key Auth"),
        ("Monitoring", "journalctl, rsyslog, /var/log analysis, Supportconfig"),
    ]

    for label, value in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0.3)
        p.paragraph_format.space_after = Pt(0.3)
        p.paragraph_format.line_spacing = Pt(9.5)
        p.paragraph_format.left_indent = Inches(0.05)
        styled_run(p, f"{label}:  ", size=Pt(8), color=DARK, bold=True)
        styled_run(p, value, size=Pt(8))

    # ═══════════════════════════════════════════════════════════
    # PROFESSIONAL EXPERIENCE
    # ═══════════════════════════════════════════════════════════
    section_heading(doc, "Professional Experience")

    p_job = doc.add_paragraph()
    p_job.paragraph_format.space_before = Pt(1.5)
    p_job.paragraph_format.space_after = Pt(0)
    p_job.paragraph_format.line_spacing = Pt(11)
    styled_run(p_job, "Linux System Administrator", size=Pt(9), color=DARK, bold=True)

    p_co = doc.add_paragraph()
    p_co.paragraph_format.space_before = Pt(0)
    p_co.paragraph_format.space_after = Pt(1.5)
    p_co.paragraph_format.line_spacing = Pt(10)
    styled_run(p_co, "Savayas Life and Balance Pvt. Ltd.", size=Pt(8), italic=True)
    styled_run(p_co, "  |  March 2024 – Present  |  Satna, MP", size=Pt(8), color=MUTED)

    exp = [
        "Administered 15+ Linux servers (SLES 15, Ubuntu 22.04) ensuring 99.5% uptime across production and staging environments for business-critical applications.",
        "Resolved 75+ P1/P2 operational incidents through systematic log analysis (/var/log/messages, journalctl, supportconfig) and root cause investigation, avg resolution under 2 hours.",
        "Provided L1/L2 technical support to 50+ end-users, resolving hardware, software, and network issues with 95% first-call resolution rate.",
        "Automated routine tasks using Bash scripts and Ansible playbooks, reducing manual effort by 40% and saving 15+ hours per week.",
        "Deployed and maintained Docker-based application stacks; managed container lifecycle, resource allocation, and network configuration.",
        "Implemented CIS-aligned server hardening: disabled legacy services, configured PAM policies, enforced firewall rules (UFW/firewalld), enabled audit logging.",
        "Maintained operational documentation: runbooks, SOPs, incident reports, and KB articles. Collaborated cross-functionally on patches and config changes with zero downtime.",
    ]
    for b in exp:
        bullet(doc, b)

    # ═══════════════════════════════════════════════════════════
    # PROJECTS
    # ═══════════════════════════════════════════════════════════
    section_heading(doc, "Projects")

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(1.5)
    p_p1.paragraph_format.space_after = Pt(0.5)
    p_p1.paragraph_format.line_spacing = Pt(10)
    styled_run(p_p1, "Automated Linux Server Hardening with Ansible", size=Pt(8.5), color=DARK, bold=True)

    bullet(doc, "Built 6 Ansible roles automating SSH hardening, firewall enforcement, audit logging, and compliance scanning across Ubuntu and openSUSE servers.")
    bullet(doc, "Enforced CIS/NIST baselines; integrated Lynis for automated compliance validation. Reduced audit prep time by 60%.")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(0.5)
    p_p2.paragraph_format.line_spacing = Pt(10)
    styled_run(p_p2, "HA Cluster P1 Incident Investigation & Root Cause Analysis", size=Pt(8.5), color=DARK, bold=True)

    bullet(doc, "Conducted full RCA on a P1 outage in a 2-node SLES HA cluster (Pacemaker/Corosync/SBD), diagnosing a 5x reboot loop caused by network isolation and fencing failure.")
    bullet(doc, "Analyzed supportconfig archives, correlated token timeouts and ARP anomalies. Delivered remediation plan; improved cluster stability by 80%.")

    # ═══════════════════════════════════════════════════════════
    # CERTIFICATIONS | EDUCATION | LANGUAGES (compact)
    # ═══════════════════════════════════════════════════════════
    section_heading(doc, "Certifications")
    certs = [
        "Linux Administration Training (CX-501) – Codenixia",
        "Linux Server Hardening & Security Training (CX-701) – Codenixia",
        "SUSE Linux Enterprise Server (SLES) Administration Training – Codenixia",
    ]
    for c in certs:
        bullet(doc, c)

    section_heading(doc, "Education")
    p_deg = doc.add_paragraph()
    p_deg.paragraph_format.space_before = Pt(1.5)
    p_deg.paragraph_format.space_after = Pt(0)
    p_deg.paragraph_format.line_spacing = Pt(10)
    styled_run(p_deg, "B.Tech – Computer Science & Engineering", size=Pt(8.5), color=DARK, bold=True)
    styled_run(p_deg, "  |  AKS University, Satna, MP", size=Pt(8), italic=True)
    styled_run(p_deg, "  |  2022–2026  |  CGPA: 8.29/10", size=Pt(8), color=MUTED)

    p_cw = doc.add_paragraph()
    p_cw.paragraph_format.space_before = Pt(0.5)
    p_cw.paragraph_format.space_after = Pt(0.5)
    p_cw.paragraph_format.line_spacing = Pt(9.5)
    p_cw.paragraph_format.left_indent = Inches(0.05)
    styled_run(p_cw, "Coursework:  ", size=Pt(8), bold=True, color=DARK)
    styled_run(p_cw, "Operating Systems, Computer Networks, DBMS, Information Security, System Admin, DSA", size=Pt(8))

    # Languages inline
    p_lang = doc.add_paragraph()
    p_lang.paragraph_format.space_before = Pt(3)
    p_lang.paragraph_format.space_after = Pt(0)
    p_lang.paragraph_format.line_spacing = Pt(10)
    styled_run(p_lang, "Languages:  ", size=Pt(8), bold=True, color=DARK)
    styled_run(p_lang, "English (Professional)  |  Hindi (Native)", size=Pt(8))

    doc.save(OUTPUT_FILE)
    print(f"Resume saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    build()
